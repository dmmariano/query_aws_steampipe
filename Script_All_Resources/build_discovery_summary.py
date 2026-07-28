#!/usr/bin/env python3
import argparse
import csv
import json
import os
import re
import sys
import traceback
from collections import Counter
from datetime import datetime, timezone
from copy import deepcopy

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


EXCLUDED_CSVS = {
    "resumo_quantidade_linhas.csv",
    "discovery_summary_by_account.csv",
    "discovery_summary_by_region.csv",
    "discovery_summary_by_resource_type.csv",
    "discovery_summary_by_resource_group.csv",
    "discovery_summary_by_account_type.csv",
    "discovery_summary_by_account_group.csv",
    "discovery_summary_by_account_category.csv",
    "discovery_summary_by_account_region.csv",
}
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
RESOURCE_CATALOG_PATH = os.path.join(SCRIPT_DIR, "resource_catalog.json")
CHUNK_SIZE = 50000
EXCEL_MAX_ROWS = 1048576
EXEC_FILL = PatternFill("solid", fgColor="164F69")
HEADER_FILL = PatternFill("solid", fgColor="EAF2F7")
HEADER_FONT = Font(bold=True, color="18202D")
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def load_resource_catalog():
    def catalog_value(value):
        if value is None:
            return ""
        value = str(value).strip()
        return "" if value.lower() in {"nan", "none", "null"} else value

    with open(RESOURCE_CATALOG_PATH, "r", encoding="utf-8") as handle:
        catalog = json.load(handle)
    groups = {
        catalog_value(group.get("id")): catalog_value(group.get("label") or group.get("id"))
        for group in catalog.get("groups", [])
        if catalog_value(group.get("id"))
    }
    resources = {}
    for table, metadata in catalog.get("resources", {}).items():
        if metadata.get("dashboard", True) is False:
            continue
        group = catalog_value(metadata.get("group")) or "other"
        resources[table] = {
            "label": catalog_value(metadata.get("label")) or table,
            "group": group,
            "group_label": groups.get(group, group),
        }
    return {
        "version": catalog.get("version", 0),
        "groups": groups,
        "resources": resources,
    }


RESOURCE_CATALOG = load_resource_catalog()


def now_iso():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def log_error(log_path, message):
    print(message)
    with open(log_path, "a", encoding="utf-8") as handle:
        handle.write(message + "\n")
        handle.write(traceback.format_exc() + "\n")


def parse_accounts_from_spc():
    spc_path = os.path.expanduser("~/.steampipe/config/aws.spc")
    accounts = {}
    if not os.path.isfile(spc_path):
        return accounts
    with open(spc_path, "r", encoding="utf-8", errors="replace") as handle:
        text = handle.read()
    for match in re.finditer(r'connection\s+"([^"]+)"\s*\{([\s\S]*?)\n\}', text):
        name = match.group(1)
        body = match.group(2)
        account_id = re.search(r'#\s*account_id\s*=\s*"([^"]+)"', body)
        account_name = re.search(r'#\s*account_name\s*=\s*"([^"]+)"', body)
        clean_id = account_id.group(1).strip() if account_id else name.replace("aws_", "")
        if re.fullmatch(r"\d{12}", clean_id):
            accounts[clean_id] = account_name.group(1).strip() if account_name else name
            accounts[f"aws_{clean_id}"] = accounts[clean_id]
    return accounts


def read_account_inventory(csv_dir, account_map, log_path):
    path = os.path.join(csv_dir, "aws_account.csv")
    accounts = {}
    if not os.path.isfile(path):
        return accounts
    try:
        for chunk in read_csv_chunks(path):
            if "account_id" not in chunk.columns:
                continue
            for _, row in chunk.iterrows():
                account_id = clean_value(row.get("account_id"))
                if not re.fullmatch(r"\d{12}", account_id):
                    continue
                connection_name = clean_value(row.get("sp_connection_name"))
                title = clean_value(row.get("title"))
                account_name = (
                    account_map.get(account_id)
                    or account_map.get(connection_name)
                    or title
                    or connection_name
                    or account_id
                )
                accounts[account_id] = account_name
                account_map[account_id] = account_name
                if connection_name:
                    account_map[connection_name] = account_name
    except pd.errors.EmptyDataError:
        return accounts
    except Exception:
        log_error(log_path, "Error reading aws_account inventory")
    return accounts


def clean_value(value):
    if value is None:
        return ""
    value = str(value).strip()
    if value.lower() in {"nan", "none", "null"}:
        return ""
    return value


def safe_filename(value):
    value = clean_value(value)
    value = re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("._")
    return value or "sem_nome"


def account_name_for(account_id, account_map):
    account_id = clean_value(account_id)
    if not account_id:
        return "Sem account_id"
    return account_map.get(account_id) or account_map.get(account_id.replace("aws_", "")) or account_id


def canonical_account_value(value, valid_accounts=None):
    value = clean_value(value)
    account_id = ""
    if re.fullmatch(r"\d{12}", value):
        account_id = value
    elif re.fullmatch(r"aws_\d{12}", value):
        account_id = value.replace("aws_", "", 1)

    if account_id and valid_accounts and account_id not in valid_accounts:
        return ""
    return account_id


def account_from_context(value, valid_accounts=None):
    value = clean_value(value)
    if not value:
        return ""
    match = re.search(r'"connection_name"\s*:\s*"aws_(\d{12})"', value)
    if not match:
        return ""
    return canonical_account_value(match.group(1), valid_accounts)


def account_series_from_chunk(chunk, valid_accounts=None):
    result = pd.Series(["Sem account_id"] * len(chunk.index), index=chunk.index)

    for column in ["account_id", "sp_connection_name", "account"]:
        if column not in chunk.columns:
            continue
        normalized = chunk[column].map(lambda value: canonical_account_value(value, valid_accounts))
        result = result.mask((result == "Sem account_id") & (normalized != ""), normalized)

    for column in ["sp_ctx", "_ctx"]:
        if column not in chunk.columns:
            continue
        normalized = chunk[column].map(lambda value: account_from_context(value, valid_accounts))
        result = result.mask((result == "Sem account_id") & (normalized != ""), normalized)

    return result


def account_category_for(account_name):
    value = clean_value(account_name).lower()
    if re.search(r"(^|[-_ ])prd($|[-_ ])|prod|production", value):
        return "Producao"
    if re.search(r"(^|[-_ ])dev($|[-_ ])|development", value):
        return "Desenvolvimento"
    if re.search(r"(^|[-_ ])hom($|[-_ ])|hml|homolog", value):
        return "Homologacao"
    if re.search(r"(^|[-_ ])qas($|[-_ ])|qa|quality", value):
        return "QAS"
    if re.search(r"sandbox|poc|lab|innovation|teste|test", value):
        return "Sandbox/POC"
    if re.search(r"backup|log-archive|shared|network|infra|security|audit", value):
        return "Plataforma"
    return "Nao classificada"


def list_resource_csvs(csv_dir):
    if not os.path.isdir(csv_dir):
        return []
    return sorted(
        name for name in os.listdir(csv_dir)
        if name.lower().endswith(".csv") and name not in EXCLUDED_CSVS
    )


def is_dashboard_resource_type(resource_type):
    return resource_type in RESOURCE_CATALOG["resources"]


def resource_metadata(resource_type):
    return RESOURCE_CATALOG["resources"].get(resource_type, {
        "label": resource_type,
        "group": "other",
        "group_label": "Other",
    })


def write_rows_csv(path, fieldnames, rows):
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})


def dataframe_from_rows(rows, columns):
    return pd.DataFrame(rows, columns=columns) if rows else pd.DataFrame(columns=columns)


def read_csv_chunks(path):
    try:
        yield from pd.read_csv(path, dtype=str, chunksize=CHUNK_SIZE, keep_default_na=False)
    except pd.errors.ParserError:
        yield from pd.read_csv(
            path,
            dtype=str,
            chunksize=CHUNK_SIZE,
            keep_default_na=False,
            engine="python",
            on_bad_lines="skip",
        )


def read_csv_for_excel(path):
    try:
        return pd.read_csv(path)
    except pd.errors.ParserError:
        return pd.read_csv(path, engine="python", on_bad_lines="skip")


def build_summary(csv_dir, summary_json, log_path):
    account_map = parse_accounts_from_spc()
    account_inventory = read_account_inventory(csv_dir, account_map, log_path)
    valid_accounts = set(account_inventory.keys())
    by_account = Counter()
    by_region = Counter()
    by_type = Counter()
    by_group = Counter()
    by_account_type = Counter()
    by_account_group = Counter()
    by_account_region = Counter()
    tables = []
    accounts_seen = set()
    accounts_seen.update(account_inventory.keys())
    regions_seen = set()
    total_resources = 0

    for filename in list_resource_csvs(csv_dir):
        path = os.path.join(csv_dir, filename)
        resource_type = filename[:-4]
        if not is_dashboard_resource_type(resource_type):
            continue
        metadata = resource_metadata(resource_type)
        resource_group = metadata["group"]
        rows = 0
        try:
            for chunk in read_csv_chunks(path):
                chunk_rows = len(chunk.index)
                rows += chunk_rows
                by_type[resource_type] += chunk_rows
                by_group[resource_group] += chunk_rows
                total_resources += chunk_rows

                account_series = account_series_from_chunk(chunk, valid_accounts)

                if "region" in chunk.columns:
                    region_series = chunk["region"].map(clean_value).replace("", "Global/sem regiao")
                else:
                    region_series = pd.Series(["Global/sem regiao"] * chunk_rows)

                for account, count in account_series.value_counts(dropna=False).items():
                    account = clean_value(account) or "Sem account_id"
                    if account == "Sem account_id":
                        continue
                    by_account[account] += int(count)
                    accounts_seen.add(account)
                    by_account_type[(account, resource_type)] += int(count)
                    by_account_group[(account, resource_group)] += int(count)
                for region, count in region_series.value_counts(dropna=False).items():
                    region = clean_value(region) or "Global/sem regiao"
                    by_region[region] += int(count)
                    regions_seen.add(region)
                for account, region in zip(account_series.tolist(), region_series.tolist()):
                    account = clean_value(account) or "Sem account_id"
                    if account == "Sem account_id":
                        continue
                    region = clean_value(region) or "Global/sem regiao"
                    by_account_region[(account, region)] += 1
        except pd.errors.EmptyDataError:
            rows = 0
        except Exception:
            log_error(log_path, f"Error summarizing {filename}")
            continue

        tables.append({
            "file": filename,
            "resource_type": resource_type,
            "resource_label": metadata["label"],
            "resource_group": metadata["group"],
            "resource_group_label": metadata["group_label"],
            "rows": rows,
        })

    by_account_rows = [
        {
            "account_id": account,
            "account_name": account_name_for(account, account_map),
            "account_category": account_category_for(account_name_for(account, account_map)),
            "count": count,
        }
        for account, count in by_account.most_common()
    ]
    counted_accounts = {item["account_id"] for item in by_account_rows}
    by_account_rows.extend([
        {
            "account_id": account,
            "account_name": account_name_for(account, account_map),
            "account_category": account_category_for(account_name_for(account, account_map)),
            "count": 0,
        }
        for account in sorted(accounts_seen - counted_accounts)
        if account != "Sem account_id"
    ])
    account_totals = {item["account_id"]: item["count"] for item in by_account_rows}
    by_account_category_counter = Counter()
    account_category_seen = {}
    for item in by_account_rows:
        by_account_category_counter[item["account_category"]] += item["count"]
        account_category_seen.setdefault(item["account_category"], set()).add(item["account_id"])
    by_account_category_rows = [
        {
            "account_category": category,
            "accounts": len(account_category_seen.get(category, set())),
            "count": count,
            "percent_total": round((count / total_resources) * 100, 2) if total_resources else 0,
        }
        for category, count in by_account_category_counter.most_common()
    ]
    by_region_rows = [
        {"region": region, "count": count}
        for region, count in by_region.most_common()
    ]
    by_type_rows = [
        {
            "resource_type": resource_type,
            "resource_label": resource_metadata(resource_type)["label"],
            "resource_group": resource_metadata(resource_type)["group"],
            "resource_group_label": resource_metadata(resource_type)["group_label"],
            "count": count,
        }
        for resource_type, count in by_type.most_common()
    ]
    by_group_rows = [
        {
            "resource_group": group,
            "resource_group_label": RESOURCE_CATALOG["groups"].get(group, group),
            "count": count,
            "percent_total": round((count / total_resources) * 100, 2) if total_resources else 0,
        }
        for group, count in by_group.most_common()
    ]
    by_account_type_rows = [
        {
            "account_id": account,
            "account_name": account_name_for(account, account_map),
            "account_category": account_category_for(account_name_for(account, account_map)),
            "resource_type": resource_type,
            "resource_label": resource_metadata(resource_type)["label"],
            "resource_group": resource_metadata(resource_type)["group"],
            "resource_group_label": resource_metadata(resource_type)["group_label"],
            "count": count,
            "percent_of_account": round((count / account_totals.get(account, count)) * 100, 2) if account_totals.get(account, 0) else 0,
            "percent_of_total": round((count / total_resources) * 100, 2) if total_resources else 0,
        }
        for (account, resource_type), count in sorted(
            by_account_type.items(),
            key=lambda item: (-item[1], item[0][0], item[0][1]),
        )
    ]
    by_account_group_rows = [
        {
            "account_id": account,
            "account_name": account_name_for(account, account_map),
            "account_category": account_category_for(account_name_for(account, account_map)),
            "resource_group": group,
            "resource_group_label": RESOURCE_CATALOG["groups"].get(group, group),
            "count": count,
            "percent_of_account": round((count / account_totals.get(account, count)) * 100, 2) if account_totals.get(account, 0) else 0,
            "percent_of_total": round((count / total_resources) * 100, 2) if total_resources else 0,
        }
        for (account, group), count in sorted(
            by_account_group.items(),
            key=lambda item: (-item[1], item[0][0], item[0][1]),
        )
    ]
    by_account_region_rows = [
        {
            "account_id": account,
            "account_name": account_name_for(account, account_map),
            "account_category": account_category_for(account_name_for(account, account_map)),
            "region": region,
            "count": count,
            "percent_of_account": round((count / account_totals.get(account, count)) * 100, 2) if account_totals.get(account, 0) else 0,
            "percent_of_total": round((count / total_resources) * 100, 2) if total_resources else 0,
        }
        for (account, region), count in sorted(
            by_account_region.items(),
            key=lambda item: (-item[1], item[0][0], item[0][1]),
        )
    ]

    summary = {
        "generatedAt": now_iso(),
        "totals": {
            "resources": total_resources,
            "accounts": len(valid_accounts) if valid_accounts else len([item for item in accounts_seen if item != "Sem account_id"]),
            "regions": len([item for item in regions_seen if item != "Global/sem regiao"]),
            "resourceTypes": len(by_type),
            "resourceGroups": len(by_group),
            "tables": len(tables),
        },
        "catalogVersion": RESOURCE_CATALOG["version"],
        "byAccount": by_account_rows,
        "byRegion": by_region_rows,
        "byType": by_type_rows,
        "byGroup": by_group_rows,
        "byAccountCategory": by_account_category_rows,
        "byAccountType": by_account_type_rows,
        "byAccountGroup": by_account_group_rows,
        "byAccountRegion": by_account_region_rows,
        "tables": sorted(tables, key=lambda item: (-item["rows"], item["resource_type"])),
    }

    os.makedirs(os.path.dirname(summary_json), exist_ok=True)
    with open(summary_json, "w", encoding="utf-8") as handle:
        json.dump(summary, handle, ensure_ascii=True, indent=2)

    write_rows_csv(
        os.path.join(csv_dir, "discovery_summary_by_account.csv"),
        ["account_id", "account_name", "account_category", "count"],
        by_account_rows,
    )
    write_rows_csv(
        os.path.join(csv_dir, "discovery_summary_by_account_category.csv"),
        ["account_category", "accounts", "count", "percent_total"],
        by_account_category_rows,
    )
    write_rows_csv(
        os.path.join(csv_dir, "discovery_summary_by_region.csv"),
        ["region", "count"],
        by_region_rows,
    )
    write_rows_csv(
        os.path.join(csv_dir, "discovery_summary_by_resource_type.csv"),
        ["resource_type", "resource_label", "resource_group", "resource_group_label", "count"],
        by_type_rows,
    )
    write_rows_csv(
        os.path.join(csv_dir, "discovery_summary_by_resource_group.csv"),
        ["resource_group", "resource_group_label", "count", "percent_total"],
        by_group_rows,
    )
    write_rows_csv(
        os.path.join(csv_dir, "discovery_summary_by_account_type.csv"),
        ["account_id", "account_name", "account_category", "resource_type", "resource_label", "resource_group", "resource_group_label", "count", "percent_of_account", "percent_of_total"],
        by_account_type_rows,
    )
    write_rows_csv(
        os.path.join(csv_dir, "discovery_summary_by_account_group.csv"),
        ["account_id", "account_name", "account_category", "resource_group", "resource_group_label", "count", "percent_of_account", "percent_of_total"],
        by_account_group_rows,
    )
    write_rows_csv(
        os.path.join(csv_dir, "discovery_summary_by_account_region.csv"),
        ["account_id", "account_name", "account_category", "region", "count", "percent_of_account", "percent_of_total"],
        by_account_region_rows,
    )
    return summary


def unique_sheet_name(base, used):
    base = (base or "sheet")[:31]
    name = base
    index = 1
    while name in used:
        suffix = f"_{index}"
        name = base[: max(0, 31 - len(suffix))] + suffix
        index += 1
    used.add(name)
    return name


def write_df(writer, used, name, df):
    sheet_name = unique_sheet_name(name, used)
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    return sheet_name


def apply_table_style(ws):
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        max_len = max(len(str(cell.value or "")) for cell in column[:80])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)


def apply_workbook_style(workbook):
    for ws in workbook.worksheets:
        apply_table_style(ws)
    if "00_executive_summary" in workbook.sheetnames:
        ws = workbook["00_executive_summary"]
        ws["A1"].fill = EXEC_FILL
        ws["A1"].font = WHITE_BOLD
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.merge_cells("A1:D1")
        for row in range(3, 7):
            ws[f"A{row}"].font = HEADER_FONT
            ws[f"B{row}"].font = Font(bold=True, size=14, color="164F69")
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["D"].width = 34
        ws.column_dimensions["E"].width = 14


def add_bar_chart(workbook, sheet_name, title, anchor, label_col=1, data_col=2, max_rows=12):
    if sheet_name not in workbook.sheetnames:
        return
    ws = workbook[sheet_name]
    rows = min(ws.max_row, max_rows + 1)
    if rows < 3:
        return
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = ""
    chart.x_axis.title = "Quantidade"
    data = Reference(ws, min_col=data_col, min_row=1, max_row=rows)
    cats = Reference(ws, min_col=label_col, min_row=2, max_row=rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 7
    chart.width = 13
    ws.add_chart(chart, anchor)


def raw_excel_tables(summary, csv_dir):
    items = []
    seen = set()
    for item in summary["tables"]:
        filename = item["file"]
        if filename in seen:
            continue
        seen.add(filename)
        items.append(item)
    if "aws_account.csv" not in seen and os.path.isfile(os.path.join(csv_dir, "aws_account.csv")):
        items.append({
            "file": "aws_account.csv",
            "resource_type": "aws_account",
            "resource_label": "AWS Account",
            "resource_group": "governance",
            "resource_group_label": "Governance",
            "rows": 0,
        })
    return items


def build_xlsx(csv_dir, xlsx_path, resumo_path, summary, log_path, account_id=None):
    print("Creating executive Excel workbook...")
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        used = set()
        totals = summary["totals"]
        executive = pd.DataFrame([
            ["Discovery AWS - Resumo executivo", "", "", "", ""],
            ["", "", "", "", ""],
            ["Total de recursos", totals["resources"], "", "Top contas por recurso", "Qtd"],
            ["Contas", totals["accounts"], "", "", ""],
            ["Regioes", totals["regions"], "", "", ""],
            ["Tipos de recurso", totals["resourceTypes"], "", "", ""],
            ["Tabelas com dados", totals["tables"], "", "", ""],
            ["Gerado em", summary["generatedAt"], "", "", ""],
        ], columns=["Indicador", "Valor", "", "Conta", "Quantidade"])
        top_accounts = summary["byAccount"][:10]
        for index, item in enumerate(top_accounts, start=2):
            if index + 1 >= len(executive):
                executive.loc[index + 1] = ["", "", "", "", ""]
            executive.at[index + 1, "Conta"] = item["account_name"]
            executive.at[index + 1, "Quantidade"] = item["count"]
        write_df(writer, used, "00_executive_summary", executive)

        write_df(writer, used, "01_by_account", dataframe_from_rows(
            summary["byAccount"], ["account_id", "account_name", "account_category", "count"]
        ))
        write_df(writer, used, "02_by_region", dataframe_from_rows(
            summary["byRegion"], ["region", "count"]
        ))
        write_df(writer, used, "03_by_resource_type", dataframe_from_rows(
            summary["byType"], ["resource_type", "resource_label", "resource_group", "resource_group_label", "count"]
        ))
        write_df(writer, used, "04_by_account_category", dataframe_from_rows(
            summary["byAccountCategory"], ["account_category", "accounts", "count", "percent_total"]
        ))
        write_df(writer, used, "05_by_account_type", dataframe_from_rows(
            summary["byAccountType"], ["account_id", "account_name", "account_category", "resource_type", "resource_label", "resource_group", "resource_group_label", "count", "percent_of_account", "percent_of_total"]
        ))
        write_df(writer, used, "06_by_account_region", dataframe_from_rows(
            summary["byAccountRegion"], ["account_id", "account_name", "account_category", "region", "count", "percent_of_account", "percent_of_total"]
        ))
        write_df(writer, used, "07_by_resource_group", dataframe_from_rows(
            summary["byGroup"], ["resource_group", "resource_group_label", "count", "percent_total"]
        ))
        write_df(writer, used, "08_by_account_group", dataframe_from_rows(
            summary["byAccountGroup"], ["account_id", "account_name", "account_category", "resource_group", "resource_group_label", "count", "percent_of_account", "percent_of_total"]
        ))
        write_df(writer, used, "09_tables", dataframe_from_rows(
            summary["tables"], ["file", "resource_type", "resource_label", "resource_group", "resource_group_label", "rows"]
        ))

        try:
            if os.path.isfile(resumo_path):
                resumo_df = pd.read_csv(resumo_path)
                write_df(writer, used, "10_raw_summary", resumo_df)
        except Exception:
            log_error(log_path, "Error adding raw summary")

        for item in raw_excel_tables(summary, csv_dir):
            filename = item["file"]
            if filename in EXCLUDED_CSVS:
                continue
            csv_path = os.path.join(csv_dir, filename)
            if not os.path.isfile(csv_path) or os.path.getsize(csv_path) == 0:
                continue
            try:
                df = read_csv_for_excel(csv_path)
                if account_id and "account_id" in df.columns:
                    df = df[df["account_id"].astype(str) == str(account_id)]
                elif account_id and "account" in df.columns:
                    df = df[df["account"].astype(str) == str(account_id)]
                elif account_id:
                    continue
                if df.empty:
                    continue
                if len(df.index) > EXCEL_MAX_ROWS - 1:
                    print(f"CSV {filename} truncated to Excel row limit.")
                    df = df.head(EXCEL_MAX_ROWS - 1)
                write_df(writer, used, filename[:-4], df)
            except pd.errors.EmptyDataError:
                print(f"Skipping empty CSV: {filename}")
            except Exception:
                log_error(log_path, f"Error processing {filename}")

        workbook = writer.book
        apply_workbook_style(workbook)
        add_bar_chart(workbook, "01_by_account", "Recursos por conta", "F2", label_col=2, data_col=4)
        add_bar_chart(workbook, "02_by_region", "Recursos por regiao", "D2", label_col=1, data_col=2)
        add_bar_chart(workbook, "03_by_resource_type", "Recursos por tipo", "G2", label_col=2, data_col=5)
        add_bar_chart(workbook, "07_by_resource_group", "Recursos por familia", "F2", label_col=2, data_col=3)


def account_summary(summary, account_id):
    account = next((item for item in summary["byAccount"] if item["account_id"] == account_id), None)
    if not account:
        return None

    by_account_type = [item for item in summary["byAccountType"] if item["account_id"] == account_id]
    by_account_group = [item for item in summary["byAccountGroup"] if item["account_id"] == account_id]
    by_account_region = [item for item in summary["byAccountRegion"] if item["account_id"] == account_id]
    table_counts = {item["resource_type"]: item["count"] for item in by_account_type}
    account_total = int(account["count"])

    scoped = deepcopy(summary)
    scoped["totals"] = {
        **summary["totals"],
        "resources": account_total,
        "accounts": 1,
        "regions": len({item["region"] for item in by_account_region if item["region"] != "Global/sem regiao"}),
        "resourceTypes": len(by_account_type),
        "resourceGroups": len(by_account_group),
        "tables": len(table_counts),
    }
    scoped["byAccount"] = [account]
    scoped["byRegion"] = [
        {"region": item["region"], "count": item["count"]}
        for item in by_account_region
    ]
    scoped["byType"] = [
        {
            "resource_type": item["resource_type"],
            "resource_label": item["resource_label"],
            "resource_group": item["resource_group"],
            "resource_group_label": item["resource_group_label"],
            "count": item["count"],
        }
        for item in by_account_type
    ]
    scoped["byGroup"] = [
        {
            "resource_group": item["resource_group"],
            "resource_group_label": item["resource_group_label"],
            "count": item["count"],
            "percent_total": item["percent_of_account"],
        }
        for item in by_account_group
    ]
    scoped["byAccountCategory"] = [{
        "account_category": account["account_category"],
        "accounts": 1,
        "count": account_total,
        "percent_total": 100,
    }]
    scoped["byAccountType"] = by_account_type
    scoped["byAccountGroup"] = by_account_group
    scoped["byAccountRegion"] = by_account_region
    scoped["tables"] = [
        {
            **item,
            "rows": table_counts[item["resource_type"]],
        }
        for item in summary["tables"]
        if item["resource_type"] in table_counts
    ]
    return scoped


def build_account_xlsx_files(csv_dir, output_dir, resumo_path, summary, log_path):
    os.makedirs(output_dir, exist_ok=True)
    for filename in os.listdir(output_dir):
        if filename.lower().endswith(".xlsx"):
            os.remove(os.path.join(output_dir, filename))

    written = []
    for account in summary["byAccount"]:
        account_id = account["account_id"]
        if not re.fullmatch(r"\d{12}", str(account_id)):
            continue
        scoped_summary = account_summary(summary, account_id)
        if not scoped_summary:
            continue
        filename = f"{safe_filename(account_id)}_{safe_filename(account['account_name'])}.xlsx"
        path = os.path.join(output_dir, filename)
        build_xlsx(csv_dir, path, resumo_path, scoped_summary, log_path, account_id=account_id)
        written.append(path)
        print(f"Account Excel saved to: {path}")
    return written


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--csv-dir", required=True)
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--summary-csv", required=True)
    parser.add_argument("--summary-json", required=True)
    parser.add_argument("--log", required=True)
    parser.add_argument("--account-xlsx-dir")
    args = parser.parse_args()

    summary = build_summary(args.csv_dir, args.summary_json, args.log)
    build_xlsx(args.csv_dir, args.xlsx, args.summary_csv, summary, args.log)
    if args.account_xlsx_dir:
        build_account_xlsx_files(args.csv_dir, args.account_xlsx_dir, args.summary_csv, summary, args.log)
    print(f"Discovery summary saved to: {args.summary_json}")
    print(f"Excel saved to: {args.xlsx}")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
        sys.exit(1)
